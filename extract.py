#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extraction des données clients BELVEDERE depuis un dossier (toutes sources).

Lit, dans le dossier choisi et ses sous-dossiers :
  - fiches clients  .doc / .docx / .txt / .rtf
  - PDF (texte, et OCR si Tesseract est installé)
  - tableaux .xlsx

Fusionne le tout par bien (Îlot + Immeuble + N°). Deux modes de fusion :
  "overwrite"  -> les nouvelles données remplacent les anciennes
  "fill"       -> on garde l'existant et on ne complète que les champs vides

Chaque client reçoit un hash de contenu et la date de modification de sa fiche,
ce qui permet de ne régénérer une lettre que si la donnée a changé.
"""

import os
import re
import sys
import json
import glob
import time
import hashlib
import subprocess

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import ocr
except ImportError:
    ocr = None

HERE = os.path.dirname(os.path.abspath(__file__))


# --------------------------------------------------------------------------- #
#  Lecture de texte — plusieurs formats, dégradation progressive              #
# --------------------------------------------------------------------------- #

def _run(cmd, timeout=40):
    try:
        out = subprocess.run(cmd, capture_output=True, timeout=timeout)
        if out.returncode == 0 and out.stdout:
            return out.stdout.decode("utf-8", errors="replace")
    except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
        pass
    return ""


def read_docx(path):
    try:
        import docx
    except ImportError:
        return ""
    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def read_doc(path):
    # 1) textutil (macOS)  2) antiword  3) catdoc  4) libreoffice --convert-to txt
    for cmd in (["textutil", "-convert", "txt", "-stdout", path],
                ["antiword", path],
                ["catdoc", path]):
        txt = _run(cmd)
        if txt.strip():
            return txt
    soffice = _find_soffice()
    if soffice:
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            _run([soffice, "--headless", "--convert-to", "txt:Text",
                  "--outdir", td, path], timeout=60)
            cand = os.path.join(td, os.path.splitext(os.path.basename(path))[0] + ".txt")
            if os.path.exists(cand):
                return open(cand, encoding="utf-8", errors="replace").read()
    return ""


def read_pdf(path):
    # couche texte d'abord
    for mod in ("pdfplumber", "fitz", "PyPDF2"):
        try:
            if mod == "pdfplumber":
                import pdfplumber
                with pdfplumber.open(path) as pdf:
                    t = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
                if t.strip():
                    return t
            elif mod == "fitz":
                import fitz
                doc = fitz.open(path)
                t = "\n".join(pg.get_text() for pg in doc)
                if t.strip():
                    return t
            elif mod == "PyPDF2":
                import PyPDF2
                r = PyPDF2.PdfReader(path)
                t = "\n".join((pg.extract_text() or "") for pg in r.pages)
                if t.strip():
                    return t
        except Exception:
            continue
    txt = _run(["pdftotext", "-layout", path, "-"])
    if txt.strip():
        return txt
    # PDF image (sans couche texte) : OCR via le moteur intégré (Apple Vision /
    # Tesseract / RapidOCR selon la machine). Indispensable pour les fiches
    # scannées en PDF (adresse, téléphone…).
    if ocr is not None:
        try:
            t = ocr.ocr_file(path, max_pages=3)
            if t and t.strip():
                return t
        except Exception:
            pass
    return _ocr_pdf(path)


def _ocr_pdf(path):
    """OCR d'un PDF scanné — seulement si Tesseract + un rasteriseur existent."""
    if not _have("tesseract"):
        return ""
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        images = []
        try:  # PyMuPDF rasterise sans dépendance externe
            import fitz
            doc = fitz.open(path)
            for i, pg in enumerate(doc):
                pix = pg.get_pixmap(dpi=300)
                img = os.path.join(td, f"p{i}.png")
                pix.save(img)
                images.append(img)
        except Exception:
            if _have("pdftoppm"):
                _run(["pdftoppm", "-r", "300", "-png", path, os.path.join(td, "p")])
                images = sorted(glob.glob(os.path.join(td, "p*.png")))
        out = []
        for img in images:
            out.append(_run(["tesseract", img, "-", "-l", "fra+eng"], timeout=120))
        return "\n".join(out)


def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return read_docx(path)
    if ext in (".doc", ".rtf"):
        return read_doc(path)
    if ext == ".txt":
        return open(path, encoding="utf-8", errors="replace").read()
    if ext == ".pdf":
        return read_pdf(path)
    return ""


def _have(binary):
    from shutil import which
    return which(binary) is not None


def _find_soffice():
    for c in ("soffice", "libreoffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if os.path.exists(c) or _have(c):
            return c
    return None


def capabilities():
    """Ce que la machine sait faire — affiché dans l'interface."""
    def ok(mod):
        try:
            __import__(mod); return True
        except ImportError:
            return False
    ocr_engine = ""
    if ocr is not None:
        try:
            ocr_engine = ocr.engine()
        except Exception:
            ocr_engine = ""
    return {
        "docx": ok("docx"),
        "xlsx": openpyxl is not None,
        "doc": _have("textutil") or _have("antiword") or _have("catdoc") or bool(_find_soffice()),
        "pdf_text": ok("pdfplumber") or ok("fitz") or ok("PyPDF2") or _have("pdftotext"),
        "pdf_ocr": bool(ocr_engine),
        "ocr_engine": ocr_engine,  # 'vision' (Apple), 'tesseract' ou ''
    }


# --------------------------------------------------------------------------- #
#  Analyse d'une fiche                                                         #
# --------------------------------------------------------------------------- #

# CIN / carte de séjour : 1-3 lettres + 4-8 chiffres, lettre finale optionnelle
# (ex. A123456, AB12345, AB10698E, CGC002533). Les passeports étrangers sont
# gérés à part.
CIN_RE = re.compile(r"\b([A-Z]{1,3}\d{4,8}[A-Z]?)\b")
DATE_RE = re.compile(r"Date\s*:?\s*(\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4})")
LABELS = ("CLIENT", "ADRESSE", "TEL", "TÉL", "C.I.N", "CIN", "QUOTE",
          "RENSEIGNEMENT", "ILOT", "ÎLOT", "IMMEUBLE", "ETAGE", "ÉTAGE",
          "PRIX", "SUPERFICIE", "RESIDENCE", "AUTORISATION", "TF N",
          "TYPE", "N°", "DESIGNATION", "AVANCE", "RESTE", "TOTAL",
          "SIGNATURE", "MONTANT")


def _clean(s):
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ")).strip()


def _is_label(line):
    up = line.upper()
    return any(up.startswith(l) or up == l for l in LABELS)


def parse_fiche(text):
    d = dict(nom_prenom="", noms=[], cin=[], quote_part=[], adresse="", tel="",
             date_fiche="", type_bien="", etage="", superficie="", prix_total="",
             avance="", reste="", note="")
    if not text:
        return d
    raw = [l for l in text.splitlines()]
    lines = [_clean(l) for l in raw]

    m = DATE_RE.search(text)
    if m:
        d["date_fiche"] = m.group(1).replace(".", "/").replace("-", "/")

    # Noms : après "CLIENT :" + lignes indentées jusqu'à ADRESSE/TEL/CIN
    noms = []
    for i, l in enumerate(lines):
        if l.upper().startswith("CLIENT"):
            after = l.split(":", 1)[1].strip() if ":" in l else ""
            if after:
                noms.append(after)
            for j in range(i + 1, min(i + 5, len(lines))):
                nxt = lines[j]
                if not nxt:
                    continue
                if nxt.upper().startswith(("ADRESSE", "TEL", "TÉL", "C.I.N", "CIN")):
                    break
                if re.search(r"[A-Za-zÀ-ÿ]", nxt) and not _is_label(nxt):
                    noms.append(nxt)
                else:
                    break
            break
    note = ""
    clean_noms = []
    for n in noms:
        mm = re.search(r"\(([^)]*repr[ée]sent[^)]*)\)", n, re.I)
        if mm:
            note = (note + " " + mm.group(0)).strip()
            n = re.sub(r"\([^)]*\)", "", n).strip()
        if n:
            clean_noms.append(n)
    d["noms"] = clean_noms
    d["nom_prenom"] = " / ".join(clean_noms)
    d["note"] = note

    # Adresse (1 ou 2 lignes)
    for i, l in enumerate(lines):
        if l.upper().startswith("ADRESSE"):
            addr = l.split(":", 1)[1].strip() if ":" in l else ""
            parts = [addr] if addr else []
            for j in range(i + 1, min(i + 3, len(lines))):
                nxt = lines[j]
                if not nxt or nxt.upper().startswith(("TEL", "TÉL", "C.I.N", "CIN")):
                    break
                if re.search(r"[A-Za-z]", nxt):
                    parts.append(nxt)
                else:
                    break
            d["adresse"] = " ".join(parts).strip()
            break

    # Téléphone
    for l in lines:
        if l.upper().startswith(("TEL", "TÉL")):
            d["tel"] = l.split(":", 1)[1].strip() if ":" in l else ""
            break
    if not d["tel"]:
        m = re.search(r"0[\s.]?[5-7](?:[\s.]?\d\d){4}", text)
        if m:
            d["tel"] = _clean(m.group(0))

    # CIN : codes lettres+chiffres, hors lignes de libellé
    cins = []
    for l in lines:
        if not l or _is_label(l):
            continue
        for mm in CIN_RE.finditer(l):
            cins.append(mm.group(1))
    seen = set()
    d["cin"] = [c for c in cins if not (c in seen or seen.add(c))]

    # Quote-part
    d["quote_part"] = re.findall(r"^\s*(100|75|66|50|33|25|20)\s*$", text, re.M)

    # Type de bien
    for key in ("Appartement", "Commerce", "Magasin", "Villa", "Bureau", "Local"):
        if re.search(key, text, re.I):
            d["type_bien"] = key
            break

    # Superficie
    m = re.search(r"(\d{2,4})\s*[mM]\s*[²2]", text)
    if m:
        d["superficie"] = m.group(1) + " m²"

    # Prix / avance / reste
    prices = re.findall(r"\d{2,3}(?:[\s.]\d{3})+(?:,\d{2})?", text)
    if prices:
        num = lambda p: float(re.sub(r"[^\d]", "", p.split(",")[0]) or 0)
        d["prix_total"] = sorted(set(prices), key=num, reverse=True)[0]
        m = re.search(r"AVANCE\s+R[EÉ]SERVATION\s*\n?\s*([\d\s.,]+)", text, re.I)
        if m:
            d["avance"] = _clean(m.group(1))
        m = re.search(r"RESTE\s+A\s+PAYER\s*\n?\s*([\d\s.,]+)", text, re.I)
        if m:
            d["reste"] = _clean(m.group(1))

    if re.search(r"\bRDC\b", text):
        d["etage"] = "RDC"
    return d


def parse_fiche_compta(text):
    """Fiche SERVICE-COMPTABILITE : mise en page différente de la notaire.

    Bloc client compact :
        Nom/Prénom
        C.I.N n° :
        QUOTE PART EN %
        <noms...>            (un ou plusieurs)
        <CIN...>             (un par acquéreur)
        <quote-parts...>     (un par acquéreur)
    """
    d = dict(nom_prenom="", noms=[], cin=[], quote_part=[], adresse="", tel="",
             date_fiche="", type_bien="", etage="", superficie="", prix_total="",
             avance="", reste="", note="")
    if not text or "COMPTABILITE" not in text.upper():
        return d
    lines = [_clean(l) for l in text.splitlines()]

    m = re.search(r"Le\s*:?\s*(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{2,4})", text)
    if m:
        d["date_fiche"] = f"{int(m.group(1)):02d}/{int(m.group(2)):02d}/{m.group(3)}"

    # Repérer l'en-tête du bloc client puis lire les lignes de données jusqu'à
    # "Renseignements Bien".
    start = None
    for i, l in enumerate(lines):
        if "QUOTE PART EN" in l.upper():
            start = i + 1
            break
    if start is not None:
        block = []
        for l in lines[start:start + 12]:
            if not l:
                continue
            if l.upper().startswith("RENSEIGNEMENT") or "RESIDENCE BELVEDERE" in l.upper():
                break
            block.append(l)
        noms, cins, qps = [], [], []
        for l in block:
            if re.fullmatch(r"(100|75|66|50|33|25|20)", l):
                qps.append(l)
            elif re.fullmatch(r"[A-Z]{1,2}\d{4,7}", l):
                cins.append(l)
            elif re.search(r"[A-Za-zÀ-ÿ]", l) and not _is_label(l):
                noms.append(l)
        d["noms"] = noms
        d["nom_prenom"] = " / ".join(noms)
        d["cin"] = cins
        d["quote_part"] = qps or (["100"] if noms else [])

    for key in ("Appartement", "Commerce", "Magasin", "Villa", "Bureau", "Local"):
        if re.search(key, text, re.I):
            d["type_bien"] = key
            break

    m = re.search(r"(\d{2,4})\s*[mM]\s*[²2]", text)
    if m:
        d["superficie"] = m.group(1) + " m²"

    # "Prix de vente après remise" = prix final ; sinon le plus grand montant
    m = re.search(r"Prix de vente apr[èe]s remise\s*\n?\s*([\d\s.]+,\d{2})", text, re.I)
    if m:
        d["prix_total"] = _clean(m.group(1))
    else:
        prices = re.findall(r"\d{2,3}(?:[\s.]\d{3})+(?:,\d{2})?", text)
        if prices:
            num = lambda p: float(re.sub(r"[^\d]", "", p.split(",")[0]) or 0)
            d["prix_total"] = sorted(set(prices), key=num, reverse=True)[0]

    if re.search(r"\bRDC\b", text):
        d["etage"] = "RDC"
    return d


def merge_fiches(primary, secondary):
    """Complète `primary` avec `secondary` (champ par champ, sans écraser)."""
    out = dict(primary)
    for k, v in secondary.items():
        cur = out.get(k)
        if k in ("noms", "cin", "quote_part"):
            if not cur and v:
                out[k] = v
        elif not cur and v:
            out[k] = v
    if not out.get("nom_prenom") and out.get("noms"):
        out["nom_prenom"] = " / ".join(out["noms"])
    return out


# --------------------------------------------------------------------------- #
#  Recherche approfondie : tous les documents d'un dossier (OCR des scans)     #
# --------------------------------------------------------------------------- #

def classify_doc(filename):
    """Devine le rôle d'un fichier d'après son nom."""
    import unicodedata
    # Normaliser et retirer les accents pour une détection fiable.
    n = unicodedata.normalize("NFKD", filename.lower())
    n = "".join(c for c in n if not unicodedata.combining(c))
    if "notaire" in n:
        return "fiche_notaire"
    if "comptabil" in n:
        return "fiche_compta"
    # CIN / carte nationale — toutes les graphies courantes (cin, cni, cnie,
    # c.n.i, carte nationale/identité, passeport, carte sejour/immatriculation).
    if ("cin" in n or "cni" in n or "cnie" in n or "c.n.i" in n
            or "identit" in n or "carte national" in n or "carte nationale" in n
            or "carte d'ident" in n or "passeport" in n or "passport" in n
            or "sejour" in n or "immatricul" in n):
        return "cin"
    if "cheq" in n or "chèq" in n or "chéq" in n:
        return "cheque"
    if "recu" in n or "reçu" in n or "reçu" in n or os.path.splitext(n)[0] == "rec":
        return "recu"
    if "vir" in n:
        return "virement"
    if "decharge" in n or "décharge" in n:
        return "decharge"
    if n.endswith((".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff")):
        return "scan"
    return "autre"


def _plausible_birthdate(s):
    """Date de naissance crédible : jj/mm/aaaa, année entre 1920 et 2012."""
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if not m:
        return False
    day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
    return 1 <= day <= 31 and 1 <= month <= 12 and 1920 <= year <= 2012


def parse_cin_card(text):
    """Extrait les données d'une carte nationale / de séjour (texte OCR).

    Gère la CIN marocaine classique (A123456), les cartes avec lettre finale
    (AB10698E) et les cartes d'immatriculation d'étranger, en s'appuyant aussi
    sur la zone de lecture automatique (MRZ : « IDMAR...<...AB10698E<... »)."""
    d = dict(cin="", naissance="", validite="", pere="", mere="", adresse="")
    if not text:
        return d
    up = text.upper()
    # 1) MRZ : la ligne « IDMAR... » contient la CIN entre chevrons «<».
    #    ex. « IDMAROPRO2T3Y<8AB10698E<<<<< » -> AB10698E
    mrz = re.search(r"IDMAR[A-Z0-9<]*?<[0-9]?([A-Z]{1,3}\d{4,8}[A-Z]?)<", up)
    if mrz:
        d["cin"] = mrz.group(1)
    # 2) « Nº AB10698E » explicite (avec lettre finale éventuelle).
    if not d["cin"]:
        m = re.search(r"N[°ºo]\s*:?\s*([A-Z]{1,3}\d{4,8}[A-Z]?)\b", up)
        if m:
            d["cin"] = m.group(1)
    # 3) repli : premier code plausible (lettre(s)+chiffres, lettre finale ok).
    if not d["cin"]:
        m = re.search(r"\b([A-Z]{1,3}\d{4,8}[A-Z]?)\b", up)
        if m:
            d["cin"] = m.group(1)
    norm = lambda s: re.sub(r"[.\-]", "/", s)
    dates = [norm(x) for x in re.findall(r"\b(\d{1,2}[.\-/]\d{1,2}[.\-/]\d{4})\b", text)]
    births = [x for x in dates if _plausible_birthdate(x)]
    if births:
        d["naissance"] = births[0]
    # Validité = une date future/récente (après la naissance retenue).
    future = [x for x in dates if not _plausible_birthdate(x)]
    if future:
        d["validite"] = future[-1]
    # Parents : motif "PRENOM ben/bent NOM" en majuscules sur la carte.
    # On isole le nom complet et on ignore les préfixes OCR bruités (Fils de…).
    pere_re = re.compile(r"([A-ZÀ-Ÿ][\wÀ-ÿ'’\-]+(?:\s+[A-ZÀ-Ÿ][\wÀ-ÿ'’\-]+)*\s+ben\s+[A-ZÀ-Ÿ][\wÀ-ÿ'’\- ]+)")
    mere_re = re.compile(r"([A-ZÀ-Ÿ][\wÀ-ÿ'’\-]+(?:\s+[A-ZÀ-Ÿ][\wÀ-ÿ'’\-]+)*\s+bent\s+[A-ZÀ-Ÿ][\wÀ-ÿ'’\- ]+)")
    for line in text.splitlines():
        l = _clean(line)
        if not d["pere"]:
            m = pere_re.search(l)
            if m and len(m.group(1)) <= 40:
                d["pere"] = m.group(1).strip()
        if not d["mere"]:
            m = mere_re.search(l)
            if m and len(m.group(1)) <= 40:
                d["mere"] = m.group(1).strip()
    # Adresse imprimée sur la carte (« Adresse : HAY … »). L'OCR perd/altère
    # souvent le mot « Adresse » (dresse, adress, messe, resse…) -> label large.
    m = re.search(r"(?i)\b[a-z]*dr?[eé]?ss?[eo]?\b\s*:?\s*([A-Z0-9][^\n]{6,70})", text)
    if not m:
        # Sur la CIN marocaine, l'adresse est souvent préfixée d'une puce
        # (« • RES LES ALLEES… », « • BAB MARZOUKA TAZA ») : signal fort.
        for line in text.splitlines():
            ls = line.strip()
            bullet = re.match(r"^[•·∙●\*\-]\s*([A-Z0-9][^\n]{5,70})", ls)
            if bullet:
                cand = bullet.group(1).strip()
                low = cand.upper()
                # on écarte les lignes MRZ / CIN / dates / libellés carte
                if ("<" not in cand and len(re.sub(r"[^A-Za-z]", "", cand)) >= 6
                        and "ROYAUME" not in low and "CARTE" not in low
                        and "NATIONALE" not in low and "VALABLE" not in low
                        and not re.match(r"^N[°ºo]", low)):
                    m = re.match(r"\s*(.+)", cand)
                    break
    if not m:
        # à défaut : une ligne commençant par un marqueur de rue marocain FORT
        # (mot complet + contenu alphabétique suffisant).
        for line in text.splitlines():
            ls = line.strip().lstrip("•·∙●*- ")
            if (re.match(r"(?i)^(HAY|RUE|LOTISSEMENT|LOT|RESIDENCE|RES|AVENUE|AV|"
                         r"BOULEVARD|BD|QUARTIER|QUART|DOUAR|BAB|DERB|ZNIQAT|ZANKAT|"
                         r"IMM|IMMEUBLE|BLOC|CITE|SIDI|OULED|AIN)\s+\S", ls)
                    and len(re.sub(r"[^A-Za-z]", "", ls)) >= 10):
                m = re.match(r"\s*(.+)", ls); break
    if m:
        a = _clean(m.group(1)).strip(" :.-")
        if (a and not re.match(r"^[0-9<]+$", a)
                and len(re.sub(r"[^A-Za-z]", "", a)) >= 6
                and "<" not in a and "ROYAUME" not in a.upper()
                and "MAROC" not in a.upper()[:6]):
            d["adresse"] = a.upper()
    return d


def parse_recu(text):
    """Numéro de reçu, montant, date depuis un reçu OCR."""
    d = dict(recu_num="", montant="", recu_date="")
    if not text:
        return d
    m = re.search(r"N[°o]\s*:?\s*(\d{3,6})", text)
    if m:
        d["recu_num"] = m.group(1)
    m = re.search(r"(\d[\d\s.]{3,})\s*dhs", text, re.I)
    if m:
        d["montant"] = _clean(m.group(1))
    m = re.search(r"\b(\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d{2,4})\b", text)
    if m:
        d["recu_date"] = _clean(m.group(1)).replace(" ", "")
    return d


# Banques marocaines fréquentes (pour repérer la banque sur chèque/virement/RIB).
_BANKS = ("ATTIJARIWAFA", "BMCE", "BANK OF AFRICA", "BANQUE POPULAIRE", "BP ",
          "CIH", "CREDIT AGRICOLE", "CREDIT DU MAROC", "SOCIETE GENERALE",
          "BMCI", "CFG", "CDG", "AL BARID", "BARID BANK", "UMNIA", "ARAB BANK")


def parse_payment(text):
    """Extrait les infos de paiement d'un chèque / virement / RIB (OCR).

    Renvoie {montant, banque, rib, payeur, mode}. Tolérant au bruit OCR."""
    d = dict(montant="", banque="", rib="", payeur="", mode="")
    if not text:
        return d
    up = text.upper()
    # Montant : « 265 500,00 » (avec DH/DHS éventuel) -> « 265500 ».
    m = re.search(r"(\d{1,3}(?:[ .]\d{3})+(?:[,.]\d{2})?)\s*(?:DH|DHS|MAD)?", up)
    if not m:
        m = re.search(r"\b(\d{4,8}(?:[,.]\d{2})?)\s*(?:DH|DHS|MAD)\b", up)
    if m:
        d["montant"] = re.sub(r"[ .](?=\d{3}\b)", "", m.group(1)).split(",")[0].replace(" ", "")
    # Mode
    if "VIREMENT" in up:
        d["mode"] = "Virement"
    elif "CHEQUE" in up or "CHÈQUE" in up:
        d["mode"] = "Chèque"
    # Banque
    for b in _BANKS:
        if b.strip() in up:
            d["banque"] = b.strip().title()
            break
    # RIB (24 chiffres) ou n° de compte long
    m = re.search(r"\b(\d{3}[ ]?\d{3}[ ]?\d{16}|\d{20,24})\b", up.replace("-", ""))
    if m:
        d["rib"] = re.sub(r"\s", "", m.group(1))
    # Donneur d'ordre / payeur (ligne après « DONNEUR D'ORDRE »)
    m = re.search(r"DONNEUR D[’'`]?ORDRE\s*[:\n]?\s*([A-ZÀ-Ÿ][A-ZÀ-Ÿ ]{4,40})", up)
    if m:
        d["payeur"] = _clean(m.group(1).title())
    return d


def scan_folder_documents(folder, root, want_ocr=True):
    """Inventorie tous les fichiers du dossier d'un bien et OCR les scans.

    Les chemins des documents sont ABSOLUS (le projet peut couvrir plusieurs
    dossiers de départ, donc on ne peut pas se fier à un seul 'root' relatif).

    Retourne (documents, enrichment) :
      documents  -> liste {name, type, rel(abs)}
      enrichment -> champs extraits des scans (CIN carte, reçu, etc.)
    """
    docs = []
    enrich = dict(naissance="", cin_validite="", pere="", mere="",
                  recu_num="", recu_montant="", recu_date="", cin_verifiee="",
                  ocr_fiche_text="", cin_src="", recu_src="", ocr_fiche_src="",
                  paiement_montant="", paiement_mode="", banque="", rib="",
                  paiement_src="", cin_adresse="", cin_adresse_src="")
    try:
        files = sorted(os.listdir(folder))
    except OSError:
        return docs, enrich

    for fn in files:
        if fn.startswith("~$") or fn.startswith("."):
            continue
        full = os.path.join(folder, fn)
        if not os.path.isfile(full):
            continue
        docs.append({"name": fn, "type": classify_doc(fn),
                     "rel": os.path.abspath(full)})

    if not (want_ocr and ocr and ocr.available()):
        return docs, enrich

    for doc in docs:
        kind, full, rel = doc["type"], os.path.join(root, doc["rel"]), doc["rel"]
        ext = os.path.splitext(full)[1].lower()
        if ext not in (".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            continue
        if kind == "cin":
            # 2 pages : recto (CIN/naissance) + verso (adresse) de la carte.
            card = parse_cin_card(ocr.ocr_file(full, max_pages=2))
            if card["cin"]:
                enrich.setdefault("cin_verifiee_list", [])
                if card["cin"] not in enrich["cin_verifiee_list"]:
                    enrich["cin_verifiee_list"].append(card["cin"])
                if not enrich["cin_verifiee"]:
                    enrich["cin_verifiee"] = card["cin"]; enrich["cin_src"] = rel
            for a, b in (("naissance", "naissance"), ("validite", "cin_validite"),
                         ("pere", "pere"), ("mere", "mere")):
                if card[a] and not enrich[b]:
                    enrich[b] = card[a]
                    if not enrich["cin_src"]:
                        enrich["cin_src"] = rel
            # Adresse imprimée sur la carte -> repli d'adresse.
            if card.get("adresse") and not enrich.get("cin_adresse"):
                enrich["cin_adresse"] = card["adresse"]
                enrich["cin_adresse_src"] = rel
        elif kind == "recu":
            r = parse_recu(ocr.ocr_file(full, max_pages=1))
            for a, b in (("recu_num", "recu_num"), ("montant", "recu_montant"),
                         ("recu_date", "recu_date")):
                if r[a] and not enrich[b]:
                    enrich[b] = r[a]; enrich["recu_src"] = rel
        elif kind in ("cheque", "virement"):
            p = parse_payment(ocr.ocr_file(full, max_pages=1))
            if p["montant"] and not enrich["paiement_montant"]:
                enrich["paiement_montant"] = p["montant"]
                enrich["paiement_mode"] = p["mode"]
                enrich["paiement_src"] = rel
            if p["banque"] and not enrich["banque"]:
                enrich["banque"] = p["banque"]
            if p["rib"] and not enrich["rib"]:
                enrich["rib"] = p["rib"]
        elif kind == "scan" and not enrich["ocr_fiche_text"]:
            t = ocr.ocr_file(full, max_pages=1)
            if "RENSEIGNEMENT" in t.upper() or "CLIENT" in t.upper():
                enrich["ocr_fiche_text"] = t; enrich["ocr_fiche_src"] = rel
    return docs, enrich


def _name_from_folder(leaf):
    """Extrait un nom de client depuis l'intitulé d'un dossier feuille.

    Ex. « A3 YAMNA BEKKAR » -> « YAMNA BEKKAR » ; « APPT 17 SOUFIANE EL AHMADI »
    -> « SOUFIANE EL AHMADI ». Retire les codes (ilot/imm/appt/étage/n°…).
    Renvoie '' si rien de plausible (au moins 2 mots alphabétiques)."""
    s = leaf
    # retirer les mots-codes de dossier / type de pièce / projet (jamais un nom)
    s = re.sub(r"(?i)\b(ilot|immeuble|imm|appartement|appart|apprt|appt|app|"
               r"commerce|magasin|local|villa|etage|rdc|desistement|desist|"
               r"bloc|rez|niveau|copie|changement|produit|"
               r"fiche|client|service|comptabilite|comptabilité|compta|notaire|"
               r"belvedere|belvédère|romana|engagement|preliminaire|préliminaire|"
               r"cin|cine|cni|carte|cheque|chèque|che|que|recu|reçu|virement|releve|relevé|"
               r"rib|bancaire|decharge|décharge|scan|attestation|contrat|acte|"
               r"compromis|passeport|reservation|réservation|dossier|"
               r"n[°oº]?)\b", " ", s)
    # retirer « A3 », « A 12 », « IMM.G », étages « 2eme »
    s = re.sub(r"(?i)\b[a-h]\s*\d+\b", " ", s)
    s = re.sub(r"(?i)\b\d+\s*(er|eme|ème|é?me)?\b", " ", s)
    s = re.sub(r"(?i)\(.*?\)", " ", s)
    s = re.sub(r"[^A-Za-zÀ-ÿ'’\- ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    toks = [t for t in s.split() if len(t) > 1]
    if len(toks) >= 2:
        return " ".join(toks).upper()
    return ""


def parse_path(path):
    """Déduit Îlot/Immeuble/N°/Étage/Type du chemin du dossier.

    Tolère de nombreuses variantes de nommage, par ex. :
      "ILOT 1/IMM E APPT 13 ETAGE 3"
      "ILOT1/APPART 21 IMM.G ILOT 1 4éme etage"
      "IMM-C COMMERCE 05", "Appartement N° 10", etc.
    """
    d = dict(ilot="", immeuble="", numero="", etage="", type_dossier="", statut="", tag="")
    parent_imm = ""   # immeuble donné par un dossier parent « IMMEUBLE X » (fiable)
    # On regarde chaque composant du chemin ; le plus profond (le bien) prime.
    for p in path.split(os.sep):
        up = p.upper().replace("É", "E").replace("È", "E")

        m = re.search(r"\bILOT\s*[:\-]?\s*(\d+)", up)
        if m:
            d["ilot"] = str(int(m.group(1)))   # « 01 » -> « 1 »

        # Dossier parent « IMMEUBLE X » (ou faute IMMEBULE/IMEUBLE) SEUL :
        # c'est la source la plus fiable du bâtiment.
        mp = re.fullmatch(r"\s*(?:IMMEUBLE|IMMEBULE|IMEUBLE)\s+([A-H])\s*", up)
        if mp:
            parent_imm = mp.group(1)

        # Immeuble : le mot « immeuble » (et ses fautes : IMMEUBLE, IMMEBULE,
        # IMEUBLE, IMMB, IMM…) suivi d'une SEULE lettre isolée (A-H).
        # Important : « IMMB G » = immeuble G (le « B » est collé au mot, ce
        # n'est PAS le bâtiment) — d'où l'exigence d'un séparateur AVANT la
        # lettre et d'une frontière de mot APRÈS.
        m = re.search(r"\b(?:IMMEUBLE|IMMEBULE|IMEUBLE|IMMB|IMM|BLOC|BL)\b[.\-:\s]+([A-H])\b", up)
        if not m:
            # repli : « IMM.A » / « IMMA » (lettre collée après un séparateur).
            m = re.search(r"\bIMM[.\-: ]?([A-H])\b", up)
        if m:
            d["immeuble"] = m.group(1)

        if "COMMERCE" in up or "MAGASIN" in up or "LOCAL" in up:
            d["type_dossier"] = "Commerce"
            m = re.search(r"(?:COMMERCE|MAGASIN|LOCAL)\s*(?:N[°O]?\s*)?0*(\d+)", up)
            if m:
                d["numero"] = m.group(1)

        if re.search(r"\bAPPART|\bAPPRT|\bAPPT|\bAPP\b|\bAPPARTEMENT", up):
            d["type_dossier"] = "Appartement"
            m = re.search(r"(?:APPARTEMENT|APPART|APPRT|APPT|APP)\s*(?:N[°O]?\s*)?0*(\d+)", up)
            if m:
                d["numero"] = m.group(1)

        if "VILLA" in up:
            d["type_dossier"] = "Villa"
            m = re.search(r"VILLA\s*(?:N[°O]?\s*)?0*(\d+)", up)
            if m:
                d["numero"] = m.group(1)

        if "RDC" in up:
            d["etage"] = "RDC"
        else:
            m = re.search(r"ETAGE\s*0*(\d+)", up) or re.search(r"\b(\d+)\s*[EÉ]ME\s*ETAGE", up) \
                or re.search(r"\b(\d+)\s*(?:ER|EME)\s*ETAGE", up)
            if m:
                d["etage"] = m.group(1)

        # Dossier commercial(e) « Vente(s) <Prénom> » -> tag agent.
        m = re.match(r"(?i)^\s*ventes?\s+(.+?)\s*$", p)
        if m:
            agent = m.group(1).strip()
            d["tag"] = agent[:1].upper() + agent[1:]

        if "DESISTEMENT" in up or "DESIST" in up:
            d["statut"] = "Désistement"

        # Format court « A02 NOM », « A7 NOM », « C15 … » : lettre d'immeuble
        # collée au n° d'appartement, souvent suivie du nom du client. On ne le
        # prend QUE si on n'a pas déjà trouvé mieux dans ce composant.
        if not d["numero"]:
            m = re.match(r"^\s*([A-H])\s?0*(\d{1,3})\b", up)
            if m:
                if not d["immeuble"]:
                    d["immeuble"] = m.group(1)
                d["numero"] = m.group(2)
                if not d["type_dossier"]:
                    d["type_dossier"] = "Appartement"

    # 0 tolérance : un dossier parent « IMMEUBLE X » explicite fait AUTORITÉ sur
    # l'immeuble (les noms de dossiers feuilles sont souvent fautifs, ex. IMMB).
    # Exception : les commerces n'ont pas d'immeuble dans le Suivi.
    if parent_imm and d["type_dossier"] != "Commerce":
        d["immeuble"] = parent_imm
    return d


# --------------------------------------------------------------------------- #
#  Lecture des tableaux xlsx                                                   #
# --------------------------------------------------------------------------- #

XLSX_FIELDS = ["ilot", "immeuble", "numerobien", "typebien", "nomprenom",
               "adresse1", "adresse2", "ville", "pays", "datecourrier",
               "commerciale", "statut", "lettreenvoyee", "observation"]


def read_xlsx_rows(path):
    rows = []
    if not openpyxl:
        return rows
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return rows
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        hdr = [str(c).strip().lower() if c is not None else "" for c in header]
        if "nomprenom" not in hdr:
            continue
        idx = {k: hdr.index(k) for k in XLSX_FIELDS if k in hdr}
        for r in it:
            g = lambda k: ("" if (idx.get(k) is None or idx[k] >= len(r)
                           or r[idx[k]] is None) else str(r[idx[k]]).strip())
            nom = g("nomprenom")
            if not nom or re.match(r"^WAFAA ZINE \d+$", nom, re.I):
                continue
            rows.append({
                "ilot": g("ilot"), "immeuble": g("immeuble"), "numero": g("numerobien"),
                "type_bien": g("typebien"), "nom_prenom": nom,
                "adresse1": g("adresse1"), "adresse2": g("adresse2"),
                "ville": g("ville"), "pays": g("pays"), "date_courrier": g("datecourrier"),
                "commerciale": g("commerciale"), "statut": g("statut"),
                "lettre_envoyee": g("lettreenvoyee"), "observation": g("observation"),
            })
    return rows


# Correspondances souples entre en-têtes Excel variés et nos champs.
_IMPORT_ALIASES = {
    "ilot": ["ilot", "îlot", "ilôt"],
    "immeuble": ["immeuble", "imm", "imm.", "bloc", "batiment", "bâtiment"],
    "numero": ["numerobien", "numero", "numéro", "n°", "no", "num", "numappt", "appartement_no"],
    "type_bien": ["typebien", "type", "produit"],
    "etage": ["etage", "étage"],
    "nom_prenom": ["nomprenom", "nom/prenom", "nom / prénom", "nom", "client", "nomprénom",
                   "nom_prenom", "nom et prenom", "nom et prénom", "acquereur", "acquéreur"],
    "cin_str": ["cin", "c.i.n", "cin n°", "cine", "cni", "carte"],
    "quote_part": ["quotepart", "quote-part", "q.p.", "qp", "quote part", "part"],
    "adresse1": ["adresse1", "adresse", "adresse 1"],
    "adresse2": ["adresse2", "adresse 2", "complement"],
    "ville": ["ville"],
    "pays": ["pays"],
    "tel": ["tel", "tél", "telephone", "téléphone", "gsm", "mobile", "contact"],
    "prix_total": ["prix", "prix total", "prixtotal", "prix (dh)", "montant"],
    "avance": ["avance", "acompte"],
    "reste": ["reste", "reste a payer", "reste à payer", "solde"],
    "date_courrier": ["date", "datecourrier", "date courrier"],
    "statut": ["statut", "etat", "état"],
    "observation": ["observation", "obs", "remarque", "note", "notes"],
    "superficie": ["superficie", "superficie (m²)", "surface", "sup", "m²", "m2"],
    "source_data": ["source données", "source donnees", "source", "provenance"],
    "dossier_source": ["dossier source", "dossier", "folder"],
    "date_vente": ["date de vente", "date vente", "date reservation", "date réservation"],
}


def _match_header(header):
    """Devine le champ-donnée correspondant à un intitulé de colonne."""
    h = str(header or "").strip().lower()
    h = re.sub(r"\s+", " ", h)
    for field, aliases in _IMPORT_ALIASES.items():
        if h in aliases:
            return field
    return None


# Libellés lisibles des champs (pour l'écran de correspondance).
FIELD_LABELS = {
    "ilot": "Îlot", "immeuble": "Immeuble", "numero": "N° du bien",
    "type_bien": "Type de bien", "etage": "Étage", "nom_prenom": "Nom / Prénom",
    "cin_str": "C.I.N", "quote_part": "Quote-part", "adresse1": "Adresse",
    "adresse2": "Adresse (compl.)", "ville": "Ville", "pays": "Pays", "tel": "Téléphone",
    "prix_total": "Prix total", "avance": "Avance", "reste": "Reste",
    "date_courrier": "Date", "statut": "Statut", "observation": "Observation",
}


def inspect_excel(path):
    """Analyse un .xlsx : pour chaque feuille, liste les colonnes, la
    correspondance auto-détectée et quelques valeurs d'exemple.

    Sert à l'écran « quelle colonne = quelle donnée » avant l'import.
    Ne modifie jamais le fichier. Renvoie un dict sérialisable JSON.
    """
    if not openpyxl:
        raise RuntimeError("Lecture Excel indisponible (openpyxl manquant).")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        if header is None:
            continue
        samples = []
        for r in it:
            samples.append(r)
            if len(samples) >= 3:
                break
        cols = []
        for i, cell in enumerate(header):
            title = "" if cell is None else str(cell).strip()
            if title == "" and all((s[i] if i < len(s) else None) in (None, "") for s in samples):
                continue  # colonne entièrement vide : on l'ignore
            vals = []
            for s in samples:
                v = s[i] if i < len(s) else None
                vals.append("" if v is None else str(v).strip())
            cols.append({"index": i, "title": title,
                         "auto": _match_header(title) or "",
                         "samples": vals})
        has_name = any(c["auto"] == "nom_prenom" for c in cols)
        sheets.append({"name": ws.title, "columns": cols,
                       "row_count": ws.max_row - 1 if ws.max_row else 0,
                       "importable": has_name})
    wb.close()
    return {"path": path, "name": os.path.basename(path), "sheets": sheets,
            "fields": [{"key": k, "label": v} for k, v in FIELD_LABELS.items()]}


def import_excel_clients(path, mapping=None):
    """Importe un .xlsx en fiches clients.

    - `mapping` (optionnel) : {nom_feuille: {champ: index_colonne}} fourni par
      l'utilisateur via l'écran de correspondance. Absent → auto-détection.
    - Enregistre pour chaque client son origine (feuille + ligne + colonnes)
      afin de pouvoir réécrire les modifications dans CE fichier plus tard.
    Ne modifie jamais le fichier source — on ne fait que lire.
    """
    if not openpyxl:
        raise RuntimeError("Lecture Excel indisponible (openpyxl manquant).")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    base = os.path.basename(path)

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        # Correspondance : utilisateur prioritaire, sinon auto.
        if mapping and ws.title in mapping:
            colmap = {f: int(i) for f, i in mapping[ws.title].items() if i not in ("", None)}
        else:
            colmap = {}
            for i, cell in enumerate(header):
                f = _match_header(cell)
                if f and f not in colmap:
                    colmap[f] = i
        if "nom_prenom" not in colmap:
            continue  # feuille sans colonne nom : ignorer
        for row_idx, r in enumerate(rows[1:], start=2):  # ligne Excel réelle (1-based, en-tête = 1)
            def g(field):
                i = colmap.get(field)
                v = r[i] if (i is not None and i < len(r)) else None
                return "" if v is None else str(v).strip()
            nom = g("nom_prenom")
            if not nom:
                continue
            noms = [n.strip() for n in re.split(r"\s*/\s*|\s*;\s*|\s*&\s*", nom) if n.strip()]
            cin_raw = g("cin_str")
            cins = [c.strip() for c in re.split(r"\s*/\s*|\s*;\s*|\s*,\s*", cin_raw) if c.strip()]
            qp_raw = g("quote_part")
            qps = [q.strip().replace("%", "") for q in re.split(r"\s*/\s*|\s*;\s*", qp_raw) if q.strip()] or (["100"] if noms else [])
            out.append(dict(
                ilot=g("ilot"), immeuble=g("immeuble"), numero=g("numero"),
                type_bien=g("type_bien") or "Appartement", etage=g("etage"),
                nom_prenom=" / ".join(noms) if noms else nom, noms=noms or [nom],
                cin=cins, cin_str=" / ".join(cins), quote_part=qps,
                adresse1=g("adresse1"), adresse2=g("adresse2"),
                ville=g("ville"), pays=g("pays"), tel=g("tel"),
                superficie=g("superficie"), prix_total=g("prix_total"), avance=g("avance"),
                reste=g("reste"), date_courrier=g("date_courrier"),
                statut=g("statut") or "Importé", observation=g("observation"),
                source_fiche="", sources=[base], source_mtime=0,
                documents=[], provenance={k: [base] for k in colmap},
                source_data=g("source_data"), dossier_source=g("dossier_source"),
                date_vente=g("date_vente"),
                # Origine Excel pour la réécriture (save-back).
                xl_sheet=ws.title, xl_row=row_idx, xl_cols=dict(colmap),
            ))
    wb.close()
    # ids + hash
    for idx, c in enumerate(out):
        c["id"] = idx + 1
        c["cin_str"] = c.get("cin_str") or " / ".join(c.get("cin", []))
        c["hash"] = _content_hash(c)
    return out


# Champs réécrits tels quels dans la cellule d'origine (valeur simple).
_SAVEBACK_FIELDS = ("ilot", "immeuble", "numero", "type_bien", "etage",
                    "nom_prenom", "cin_str", "adresse1", "adresse2", "ville",
                    "pays", "tel", "prix_total", "avance", "reste",
                    "date_courrier", "statut", "observation")


def save_excel_back(path, clients, backup_dir=None):
    """Réécrit les valeurs (éventuellement modifiées) des clients dans LE
    fichier Excel d'origine, après avoir créé une sauvegarde horodatée.

    - Seules les colonnes qui avaient été reconnues (mappées) sont réécrites ;
      les autres colonnes et toute la mise en forme sont préservées.
    - `clients` doit porter xl_sheet / xl_row / xl_cols (posés à l'import).
    Renvoie {ok, backup, written, skipped}.
    """
    if not openpyxl:
        raise RuntimeError("Écriture Excel indisponible (openpyxl manquant).")
    if not os.path.isfile(path):
        raise FileNotFoundError(path)

    backup = _backup_excel(path, backup_dir)
    # On ouvre SANS read_only et SANS data_only pour conserver formules/format.
    wb = openpyxl.load_workbook(path)
    written = skipped = 0
    for c in clients:
        sheet = c.get("xl_sheet")
        row = c.get("xl_row")
        cols = c.get("xl_cols") or {}
        if not sheet or not row or sheet not in wb.sheetnames:
            skipped += 1
            continue
        ws = wb[sheet]
        for field, col_idx in cols.items():
            if field not in _SAVEBACK_FIELDS:
                continue
            val = c.get(field, "")
            if isinstance(val, (list, tuple)):
                val = " / ".join(str(x) for x in val)
            ws.cell(row=int(row), column=int(col_idx) + 1, value=val)
        written += 1
    wb.save(path)
    return {"ok": True, "backup": os.path.basename(backup) if backup else "",
            "written": written, "skipped": skipped}


def _backup_excel(path, backup_dir=None):
    """Copie le .xlsx dans un sous-dossier Sauvegardes/ horodaté."""
    import shutil, datetime
    folder = backup_dir or os.path.join(os.path.dirname(path), "Sauvegardes")
    os.makedirs(folder, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%Hh%M")
    root, ext = os.path.splitext(os.path.basename(path))
    dest = os.path.join(folder, f"{root}_{stamp}{ext}")
    n = 2
    while os.path.exists(dest):
        dest = os.path.join(folder, f"{root}_{stamp}_{n}{ext}")
        n += 1
    shutil.copy2(path, dest)
    return dest


def list_excel_backups(path, backup_dir=None):
    """Liste les sauvegardes existantes d'un fichier Excel (plus récent d'abord)."""
    folder = backup_dir or os.path.join(os.path.dirname(path), "Sauvegardes")
    if not os.path.isdir(folder):
        return []
    root = os.path.splitext(os.path.basename(path))[0]
    items = []
    for f in os.listdir(folder):
        if f.startswith(root) and f.lower().endswith((".xlsx", ".xls")):
            full = os.path.join(folder, f)
            try:
                st = os.stat(full)
                items.append({"name": f, "path": full, "size": st.st_size,
                              "mtime": st.st_mtime})
            except OSError:
                pass
    items.sort(key=lambda x: x["mtime"], reverse=True)
    return items


# --------------------------------------------------------------------------- #
#  Fusion                                                                      #
# --------------------------------------------------------------------------- #

FIELDS = ("ilot", "immeuble", "numero", "type_bien", "etage", "nom_prenom",
          "cin_str", "quote_part", "adresse1", "adresse2", "ville", "pays",
          "tel", "superficie", "prix_total", "avance", "reste", "date_courrier",
          "statut", "observation")


def _key(ilot, immeuble, numero, type_bien=""):
    return f"{ilot}|{immeuble}|{numero}|{(type_bien or '').lower()[:4]}"


def _content_hash(c):
    basis = "|".join(str(c.get(f, "")) for f in FIELDS)
    return hashlib.md5(basis.encode("utf-8")).hexdigest()[:12]


DOC_EXTS = ("*.doc", "*.docx", "*.txt", "*.rtf", "*.pdf")


def scan(roots, mode="overwrite", existing=None, deep=True, progress=None):
    """Scanne un ou plusieurs dossiers. `roots` peut être un chemin ou une liste.

    mode: overwrite (remplace) | fill (complète l'existant / ajout de dossier).
    deep: OCR des scans (CIN, reçus…).
    """
    emit = progress or (lambda *_: None)
    if isinstance(roots, str):
        roots = [roots]
    roots = [os.path.abspath(os.path.expanduser(r)) for r in roots if r]

    prev = {}
    if existing:
        for c in existing:
            prev[_key(c["ilot"], c["immeuble"], c["numero"], c["type_bien"])] = c

    clients = {}

    skip = os.path.abspath(HERE)  # ne pas scanner le dossier de l'outil lui-même
    inside_tool = lambda p: os.path.abspath(p).startswith(skip)

    want_ocr = deep and ocr is not None and ocr.available()

    # Tout fichier exploitable sert à repérer un dossier de bien : fiches, mais
    # aussi scans, CIN, reçus… Un dossier sans .doc mais avec des scans compte.
    bien_exts = DOC_EXTS + ("*.jpg", "*.jpeg", "*.png", "*.tif", "*.tiff")
    all_files = []
    for root in roots:
        for ext in bien_exts:
            all_files += glob.glob(os.path.join(root, "**", ext), recursive=True)
    all_files = [f for f in all_files
                 if "~$" not in os.path.basename(f) and not inside_tool(f)]

    # Un dossier est un "bien" si son chemin décrit un bien (ILOT/IMM/…)
    folders = {}
    for f in all_files:
        folder = os.path.dirname(f)
        pinfo = parse_path(folder)
        if pinfo["numero"] or pinfo["ilot"]:
            folders.setdefault(folder, True)
    items = sorted(folders)
    emit(f"{len(items)} dossier(s) de bien trouvé(s)"
         + (" — OCR actif" if want_ocr else ""), 3)

    # Pré-OCR en parallèle de tous les scans (CIN, reçus…) pour remplir le cache.
    if want_ocr:
        ocr_targets = []
        for folder in items:
            try:
                for x in os.listdir(folder):
                    if x.startswith((".", "~$")):
                        continue
                    kind = classify_doc(x)
                    if kind in ("cin", "recu", "scan") and \
                            x.lower().endswith((".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff")):
                        ocr_targets.append(os.path.join(folder, x))
            except OSError:
                pass
        if ocr_targets:
            def ocr_prog(done, total):
                emit(f"Lecture des scans (OCR) : {done}/{total}",
                     3 + int(2 * done / max(1, total)))
            emit(f"Lecture des scans (OCR) : 0/{len(ocr_targets)}", 3)
            ocr.warm_cache(ocr_targets, max_pages=1, progress=ocr_prog)

    def _within(path):
        ap = os.path.abspath(path)
        for r in roots:
            if ap == r or ap.startswith(r + os.sep):
                return os.path.relpath(ap, r)
        return os.path.basename(ap)

    for i, folder in enumerate(items):
        emit(f"Analyse : {_within(folder)}", 3 + int(57 * (i + 1) / max(1, len(items))))
        pinfo = parse_path(folder)

        try:
            flist = [os.path.join(folder, x) for x in os.listdir(folder)
                     if os.path.isfile(os.path.join(folder, x))
                     and not x.startswith("~$")]
        except OSError:
            flist = []
        # Choisir la fiche de DONNÉES, pas la version SIGNÉE/scannée. La fiche
        # « FICHE CLIENT …-notaire.doc » contient les champs (adresse, tél…) ;
        # une « fiche notaire signée.pdf » n'est que la page de signature.
        # On privilégie : .doc/.docx « FICHE CLIENT », puis tout .doc/.docx,
        # puis à défaut le reste. On évite « changement de produit » (ancien).
        def _pick_fiche(keyword):
            cands = [x for x in flist if keyword in os.path.basename(x).lower()]
            if not cands:
                return None
            def rank(p):
                b = os.path.basename(p).lower()
                ext = os.path.splitext(b)[1]
                r = 0
                if "fiche client" in b or "fiche-client" in b:
                    r -= 4
                if ext in (".doc", ".docx"):
                    r -= 3                     # fiche typée (données) >> PDF signé
                if "sign" in b or "scan" in b or ext == ".pdf":
                    r += 2                     # version signée/scannée : moins bien
                if "changement" in b or "produit" in b or "copie" in b:
                    r += 3                     # ancienne fiche : en dernier
                return r
            return sorted(cands, key=rank)[0]
        notaire = _pick_fiche("notaire")
        compta = _pick_fiche("comptabil")
        # Chemins ABSOLUS pour la provenance (projet multi-dossiers).
        relof = lambda p: os.path.abspath(p) if p else ""

        # Parser chaque source séparément pour tracer la provenance.
        f_not = parse_fiche(read_any(notaire)) if notaire else {}
        f_com = parse_fiche_compta(read_any(compta)) if compta else {}

        documents, enrich = scan_folder_documents(folder, folder, want_ocr=want_ocr)

        f_ocr = {}
        ocr_src = ""
        if not (f_not.get("nom_prenom") or f_com.get("nom_prenom")) and enrich.get("ocr_fiche_text"):
            f_ocr = parse_fiche(enrich["ocr_fiche_text"])
            ocr_src = enrich.get("ocr_fiche_src", "")

        # prov : pour chaque champ, liste des fichiers (rel) qui le fournissent.
        prov = {}

        def pick(field, candidates):
            """candidates = [(valeur, source_rel), ...] par ordre de priorité.
            Retourne la 1re valeur non vide et enregistre toutes les sources
            qui portent CETTE valeur (pour l'info « d'où vient la donnée »)."""
            chosen = ""
            for val, src in candidates:
                if val:
                    chosen = val
                    break
            if chosen:
                srcs = []
                norm = lambda x: re.sub(r"\s+", " ", str(x)).strip().lower()
                for val, src in candidates:
                    if val and src and norm(val) == norm(chosen) and src not in srcs:
                        srcs.append(src)
                if srcs:
                    prov[field] = srcs
            return chosen

        src_not, src_com = relof(notaire), relof(compta)
        nom = pick("nom_prenom", [(f_not.get("nom_prenom"), src_not),
                                  (f_com.get("nom_prenom"), src_com),
                                  (f_ocr.get("nom_prenom"), ocr_src)])
        noms = f_not.get("noms") or f_com.get("noms") or f_ocr.get("noms") or []
        raw_cins = f_not.get("cin") or f_com.get("cin") or f_ocr.get("cin") or []
        # Toutes les CIN confirmées par une carte scannée (liste).
        verified = list(enrich.get("cin_verifiee_list") or
                        ([enrich["cin_verifiee"]] if enrich.get("cin_verifiee") else []))

        # Nombre d'acquéreurs attendu (quote-part ou noms).
        n_owners = max(len(noms), len([q for q in (f_not.get("quote_part")
                       or f_com.get("quote_part") or []) if q]), 1)

        # Réconcilier : les CIN VÉRIFIÉES (carte d'identité) sont prioritaires et
        # placées en premier ; on complète ensuite avec les CIN des fiches, sans
        # dépasser le nombre d'acquéreurs (pour éviter le bruit OCR).
        ordered = []
        for v in verified:
            if v and v not in ordered:
                ordered.append(v)
        for c in raw_cins:
            if c and c not in ordered:
                ordered.append(c)
        # On garde au plus n_owners CIN (les vérifiées d'abord). Les éventuelles
        # CIN en trop sont conservées separement comme "autres codes détectés".
        cins = ordered[:n_owners] if ordered else []
        extras = ordered[n_owners:]
        cin_verified_set = [c for c in cins if c in verified]
        cin_verif = verified[0] if verified else ""

        if cins:
            prov["cin"] = [s for s, dd in ((src_not, f_not), (src_com, f_com), (ocr_src, f_ocr))
                           if s and dd.get("cin")]
            if enrich.get("cin_src") and verified:
                prov.setdefault("cin", [])
                if enrich["cin_src"] not in prov["cin"]:
                    prov["cin"].append(enrich["cin_src"])

        qp = pick("quote_part", [(f_not.get("quote_part"), src_not),
                                 (f_com.get("quote_part"), src_com)]) or \
             f_not.get("quote_part") or f_com.get("quote_part") or []
        adresse = pick("adresse1", [(f_not.get("adresse"), src_not),
                                    (f_ocr.get("adresse"), ocr_src),
                                    (f_com.get("adresse"), src_com),
                                    (enrich.get("cin_adresse"),
                                     enrich.get("cin_adresse_src"))])
        tel = pick("tel", [(f_not.get("tel"), src_not), (f_ocr.get("tel"), ocr_src)])
        superficie = pick("superficie", [(f_not.get("superficie"), src_not),
                                         (f_com.get("superficie"), src_com)])
        prix = pick("prix_total", [(f_not.get("prix_total"), src_not),
                                   (f_com.get("prix_total"), src_com)])
        avance = pick("avance", [(f_not.get("avance"), src_not),
                                 (enrich.get("recu_montant"), enrich.get("recu_src"))])
        reste = pick("reste", [(f_not.get("reste"), src_not)])
        date_fiche = pick("date_courrier", [(f_not.get("date_fiche"), src_not),
                                            (f_com.get("date_fiche"), src_com)])
        etage = pick("etage", [(f_not.get("etage"), src_not), (f_com.get("etage"), src_com)]) or pinfo["etage"]
        if etage and etage == pinfo["etage"] and "etage" not in prov:
            prov["etage"] = ["(dossier)"]

        if not noms and nom:
            noms = [nom]
        if not nom and noms:
            nom = " / ".join(noms)

        # REPLI 1 : aucun nom via les fiches -> NOM DU DOSSIER
        # (ex. « A3 YAMNA BEKKAR », « APPT 17 SOUFIANE EL AHMADI »).
        if not nom:
            fn = _name_from_folder(os.path.basename(folder))
            if fn:
                nom = fn
                prov["nom_prenom"] = ["(dossier)"]
        # REPLI 2 : nom dans un NOM DE FICHIER (ex. « Cin YAACOUBI OTHMANE.pdf »,
        # « ENGAGEMENT PRELIMINAIRE yamna bekkar.docx »).
        if not nom:
            for x in flist:
                base_x = os.path.splitext(os.path.basename(x))[0]
                # retirer le mot-clé type de pièce en préfixe
                base_x = re.sub(r"(?i)^\s*(cin|cine|cni|carte|fiche|re[çc]u|recu|"
                                r"ch[eè]que|virement|engagement|preliminaire|"
                                r"compta\w*|notaire|scan|copie|\d+)\b[\s.\-_]*",
                                " ", base_x)
                fn2 = _name_from_folder(base_x)
                if fn2:
                    nom = fn2
                    prov["nom_prenom"] = ["(nom de fichier)"]
                    break
        if nom and not noms:
            noms = [n.strip() for n in re.split(r"\s+ET\s+|\s*/\s*|\s*&\s*",
                                                nom, flags=re.I) if n.strip()]

        if not nom and not pinfo["numero"]:
            continue

        # Données issues des scans (provenance = la pièce concernée)
        for fld, srckey in (("naissance", "cin_src"), ("cin_validite", "cin_src"),
                            ("pere", "cin_src"), ("mere", "cin_src"),
                            ("recu_num", "recu_src"), ("recu_montant", "recu_src"),
                            ("recu_date", "recu_src"), ("cin_verifiee", "cin_src")):
            if enrich.get(fld) and enrich.get(srckey):
                prov[fld] = [enrich[srckey]]

        type_bien = pinfo["type_dossier"] or f_not.get("type_bien") or f_com.get("type_bien") or "Appartement"
        prov["type_bien"] = ["(dossier)"]
        prov["ilot"] = prov["immeuble"] = prov["numero"] = ["(dossier)"]
        used = [s for s in (src_not, src_com) if s]
        mtime = max([os.path.getmtime(x) for x in flist] or [0])
        rec = dict(
            ilot=pinfo["ilot"], immeuble=pinfo["immeuble"], numero=pinfo["numero"],
            type_bien=type_bien, etage=etage,
            nom_prenom=nom, noms=noms, cin=cins, cin_str=" / ".join(cins),
            quote_part=qp, adresse1=adresse, adresse2="", ville="", pays="",
            tel=tel, superficie=superficie, prix_total=prix, avance=avance,
            reste=reste, date_courrier=date_fiche,
            statut=pinfo["statut"] or ("À compléter" if not nom else "Vendu"),
            observation=f_not.get("note", "") or f_com.get("note", ""),
            source_fiche=used[0] if used else "", sources=used, source_mtime=mtime,
            documents=documents,
            naissance=enrich.get("naissance", ""),
            cin_validite=enrich.get("cin_validite", ""),
            pere=enrich.get("pere", ""), mere=enrich.get("mere", ""),
            recu_num=enrich.get("recu_num", ""),
            recu_montant=enrich.get("recu_montant", ""),
            recu_date=enrich.get("recu_date", ""),
            paiement_montant=enrich.get("paiement_montant", ""),
            paiement_mode=enrich.get("paiement_mode", ""),
            banque=enrich.get("banque", ""),
            rib=enrich.get("rib", ""),
            cin_verifiee=cin_verif,
            cin_verified=cin_verified_set,   # CIN confirmées par carte d'identité
            cin_extras=extras,               # autres codes détectés (à vérifier)
            provenance=prov,
        )
        rec["source_folder"] = folder
        rec["tag"] = pinfo.get("tag", "")        # agent/commercial(e) du dossier
        k = _key(rec["ilot"], rec["immeuble"], rec["numero"], type_bien)
        if k in clients and clients[k].get("source_folder") != folder:
            # Même îlot+imm+n°+type mais AUTRE dossier -> doublon : on garde les deux.
            rec["is_duplicate"] = True
            clients[k]["is_duplicate"] = True
            clients[k + "#DUP#" + os.path.basename(folder)] = rec
        else:
            clients[k] = rec

    xlsx_files = []
    for root in roots:
        xlsx_files += glob.glob(os.path.join(root, "**", "*.xlsx"), recursive=True)
    xlsx_files = [f for f in xlsx_files
                  if "~$" not in os.path.basename(f) and not inside_tool(f)]
    emit(f"{len(xlsx_files)} fichier(s) Excel trouvé(s)", 62)
    for i, f in enumerate(xlsx_files):
        rel = _within(f)
        emit(f"Excel : {rel}", 62 + int(20 * (i + 1) / max(1, len(xlsx_files))))
        for row in read_xlsx_rows(f):
            k = _key(row["ilot"], row["immeuble"], row["numero"], row["type_bien"])
            if k in clients:
                rec = clients[k]
                for fld in ("adresse1", "adresse2", "ville", "pays", "commerciale",
                            "lettre_envoyee", "observation", "date_courrier"):
                    if not rec.get(fld) and row.get(fld):
                        rec[fld] = row[fld]
                if rel not in rec["sources"]:
                    rec["sources"].append(rel)
            else:
                clients[k] = dict(
                    ilot=row["ilot"], immeuble=row["immeuble"], numero=row["numero"],
                    type_bien=row["type_bien"] or "Appartement", etage="",
                    nom_prenom=row["nom_prenom"], noms=[row["nom_prenom"]], cin=[],
                    cin_str="", quote_part=["100"], adresse1=row["adresse1"],
                    adresse2=row["adresse2"], ville=row["ville"], pays=row["pays"],
                    tel="", superficie="", prix_total="", avance="", reste="",
                    date_courrier=row.get("date_courrier", ""),
                    statut=row.get("statut", ""), observation=row.get("observation", ""),
                    source_fiche="", sources=[rel], source_mtime=os.path.getmtime(f),
                )

    emit("Fusion des données...", 86)
    # Appliquer le mode de fusion contre l'existant
    if existing and mode == "fill":
        for k, old in prev.items():
            if k not in clients:
                clients[k] = old
            else:
                new = clients[k]
                for fld in FIELDS:
                    if (not new.get(fld)) and old.get(fld):
                        new[fld] = old[fld]
                new["edited"] = old.get("edited", False)

    def sortkey(c):
        return (c["ilot"] or "9", c["immeuble"] or "Z",
                int(re.sub(r"\D", "", c["numero"]) or 0))
    result = sorted(clients.values(), key=sortkey)
    for idx, c in enumerate(result):
        c["id"] = idx + 1
        c["cin_str"] = c.get("cin_str") or " / ".join(c.get("cin", []))
        c["hash"] = _content_hash(c)
    emit(f"Terminé : {len(result)} clients.", 100)
    return result


# --------------------------------------------------------------------------- #
#  Export Excel                                                                #
# --------------------------------------------------------------------------- #

def export_xlsx(clients, out_path):
    if not openpyxl:
        return False
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients BELVEDERE"
    cols = ["ID", "Îlot", "Immeuble", "N°", "Type", "Étage", "Nom / Prénom",
            "C.I.N", "Quote-part", "Adresse", "Ville", "Tél", "Superficie",
            "Prix total (DH)", "Avance", "Reste", "Date", "Statut", "Source"]
    ws.append(cols)
    head_fill = PatternFill("solid", fgColor="143A82")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for cl in clients:
        ws.append([cl["id"], cl["ilot"], cl["immeuble"], cl["numero"], cl["type_bien"],
                   cl["etage"], cl["nom_prenom"], cl["cin_str"],
                   " / ".join(cl["quote_part"]), cl["adresse1"], cl["ville"], cl["tel"],
                   cl["superficie"], cl["prix_total"], cl["avance"], cl["reste"],
                   cl["date_courrier"], cl["statut"], cl.get("source_fiche", "")])
    for i, w in enumerate([5, 6, 9, 5, 12, 7, 28, 16, 10, 40, 12, 16, 11, 16, 14, 14, 12, 14, 30], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)
    return True


def main():
    import datetime
    root = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.dirname(HERE)
    print(f"[*] Dossier : {root}")
    print(f"[*] Capacités : {capabilities()}")
    clients = scan(root, progress=lambda m, p: print(f"[{p:3d}%] {m}"))
    data = {"root": root, "generated": datetime.date.today().isoformat(), "clients": clients}
    json.dump(data, open(os.path.join(HERE, "data.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    export_xlsx(clients, os.path.join(HERE, "BELVEDERE_clients.xlsx"))
    print(f"[+] {len(clients)} clients -> data.json + BELVEDERE_clients.xlsx")


if __name__ == "__main__":
    main()

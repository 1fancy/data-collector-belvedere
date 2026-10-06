#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génération des lettres de publipostage BELVEDERE.

Les lettres ne sont plus générées en masse : une lettre est produite à la
demande (clic dans l'interface, ou génération automatique si l'option est
cochée). Le résultat est mis en cache dans Letters/<id_nom_cin>/ et n'est
régénéré que si le contenu du client a changé (comparaison d'empreinte).
"""

import os
import re
import json

try:
    import docx
except ImportError:
    docx = None

HERE = os.path.dirname(os.path.abspath(__file__))

MERGE_MAP = {
    "NomPrenom":   lambda c: c.get("nom_prenom", ""),
    "Adresse1":    lambda c: c.get("adresse1", ""),
    "Adresse2":    lambda c: c.get("adresse2", ""),
    "Ville":       lambda c: c.get("ville", ""),
    "Pays":        lambda c: c.get("pays", ""),
    "DateCourrier": lambda c: c.get("date_courrier", ""),
    "TypeBien":    lambda c: c.get("type_bien", ""),
    "Ilot":        lambda c: c.get("ilot", ""),
    "Immeuble":    lambda c: c.get("immeuble", ""),
    "NumeroBien":  lambda c: c.get("numero", ""),
}

FIELD_RE = re.compile(r"[«<]{1,2}\s*([A-Za-z0-9_]+)\s*[»>]{1,2}")


def _safe(s):
    s = re.sub(r"[^\w\s-]", "", s or "", flags=re.U).strip()
    return re.sub(r"\s+", "_", s)[:50]


def folder_name(c):
    nom = _safe(c.get("nom_prenom", "").split(" / ")[0]) or "CLIENT"
    cin = _safe((c.get("cin_str") or "").split(" / ")[0])
    return f"{c['id']:03d}_{nom}" + (f"_{cin}" if cin else "")


def find_template(root):
    for f in os.listdir(root):
        if "lettre" in f.lower() and f.lower().endswith(".docx"):
            return os.path.join(root, f)
    return None


# --------------------------------------------------------------------------- #
#  Remplissage du .docx (champs pouvant être découpés en plusieurs runs)      #
# --------------------------------------------------------------------------- #

def _fill_paragraph(par, values):
    full = "".join(r.text for r in par.runs)
    if "«" not in full and "<<" not in full:
        return
    new = FIELD_RE.sub(lambda m: str(values.get(m.group(1), m.group(0))), full)
    if new != full and par.runs:
        par.runs[0].text = new
        for r in par.runs[1:]:
            r.text = ""


def fill_docx(template, client, out_path):
    if not docx or not template or not os.path.exists(template):
        return False
    values = {k: fn(client) for k, fn in MERGE_MAP.items()}
    d = docx.Document(template)
    for p in d.paragraphs:
        _fill_paragraph(p, values)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fill_paragraph(p, values)
    d.save(out_path)
    return True


# --------------------------------------------------------------------------- #
#  Version HTML (aperçu web + impression PDF)                                 #
# --------------------------------------------------------------------------- #

BODY = """\
<p>Objet : <strong>Disponibilité du titre foncier – Résidence Belvédère</strong></p>
<p>Madame, Monsieur,</p>
<p>Nous vous informons que le titre foncier du bien que vous avez réservé à la
Résidence BELVÉDÈRE, projet de la société CORAIL DEVELOPPEMENT, est désormais
disponible.</p>
<p>Par conséquent, nous vous invitons à vous présenter, dans un délai de deux
semaines à compter de la réception de la présente lettre, à notre bureau de vente
situé avenue Lalla Meryem à Sala Al Jadida, afin de régulariser la situation de
votre dossier et de finaliser les démarches administratives nécessaires. Vous serez
ensuite orienté vers le notaire chargé d'établir votre contrat d'acquisition.</p>
<p>Nous vous prions de bien vouloir accomplir ces démarches dans le délai indiqué
afin d'éviter tout désagrément lié à un éventuel retard.</p>
<p>Nous vous remercions de votre confiance et restons à votre disposition pour tout
complément d'information aux numéros suivants : 06 66 88 66 20 ou 06 66 33 09 13.</p>
<p>Veuillez agréer, Madame, Monsieur, l'expression de nos salutations distinguées.</p>
<p class="sign">La Direction Commerciale</p>
"""


def build_html(client, logo_data_uri="", company=None, date_override="",
               show_print=True):
    """Construit le HTML de la lettre (chaîne) — sert au fichier, à l'aperçu et au PDF."""
    company = company or {}
    v = {k: fn(client) for k, fn in MERGE_MAP.items()}
    date_txt = date_override or v["DateCourrier"] or "........................"
    addr = "<br>".join(x for x in (v["Adresse1"], v["Adresse2"], v["Ville"], v["Pays"]) if x)
    ref = " &nbsp;·&nbsp; ".join(x for x in (
        f"Réf. : {v['TypeBien']}" if v["TypeBien"] else "",
        f"Îlot {v['Ilot']}" if v["Ilot"] else "",
        f"Immeuble {v['Immeuble']}" if v["Immeuble"] else "",
        f"N° {v['NumeroBien']}" if v["NumeroBien"] else "") if x)
    logo = f'<img src="{logo_data_uri}" style="height:70px">' if logo_data_uri else ""

    cname = (company.get("name") or "").strip()
    header = f'<div class="coname">{cname}</div>' if cname else ""
    footer_bits = [company.get("address", ""), company.get("phone", ""),
                   company.get("email", ""), company.get("extra", "")]
    footer = " &nbsp;·&nbsp; ".join(x for x in footer_bits if x and x.strip())
    footer_html = f'<div class="foot">{footer}</div>' if footer else ""

    print_btn = ('<div class="noprint"><button onclick="window.print()">'
                 'Imprimer / Enregistrer en PDF</button></div>') if show_print else ""

    return f"""<!doctype html><html lang="fr"><head><meta charset="utf-8">
<title>Lettre – {v['NomPrenom']}</title><style>
 @page {{ size:A4; margin:2.2cm; }}
 body {{ font-family:"Times New Roman",Georgia,serif; font-size:13pt; line-height:1.55;
        color:#111; max-width:18cm; margin:1cm auto; padding:1.4cm; background:#fff; }}
 .top {{ display:flex; justify-content:space-between; align-items:flex-start; margin-bottom:1cm; }}
 .brand {{ display:flex; flex-direction:column; gap:4px; }}
 .coname {{ font-weight:bold; color:#143A82; font-size:12pt; }}
 .addr {{ width:9cm; border:1px solid #ccc; padding:.5cm .7cm; border-radius:4px; align-self:flex-end; }}
 .date {{ text-align:right; margin:.6cm 0 1cm; }}
 .ref {{ font-weight:bold; margin-bottom:1cm; }}
 p {{ text-align:justify; margin:.5cm 0; }}
 .sign {{ margin-top:1.4cm; font-weight:bold; text-align:right; }}
 .foot {{ margin-top:1.6cm; padding-top:.4cm; border-top:1px solid #ccc; font-size:9.5pt;
          color:#555; text-align:center; }}
 .noprint {{ text-align:center; margin-bottom:1cm; }}
 .noprint button {{ padding:.55em 1.3em; font-size:12pt; cursor:pointer; background:#143A82;
        color:#fff; border:0; border-radius:6px; font-weight:600; }}
 @media print {{ body{{margin:0;padding:0;max-width:none}} .noprint{{display:none}} }}
</style></head><body>
{print_btn}
<div class="top"><div class="brand">{logo}{header}</div><div class="addr">{addr or '&nbsp;'}</div></div>
<div class="date">Rabat, le {date_txt}</div>
<div class="ref">{ref}</div>
{BODY}{footer_html}</body></html>"""


def fill_html(client, out_path, logo_data_uri="", company=None, date_override=""):
    html = build_html(client, logo_data_uri, company, date_override)
    open(out_path, "w", encoding="utf-8").write(html)
    return True


# --------------------------------------------------------------------------- #
#  Cache                                                                        #
# --------------------------------------------------------------------------- #

def _meta_path(fol):
    return os.path.join(fol, ".meta.json")


def is_cached(client, fol):
    """La lettre en cache correspond-elle encore au contenu actuel du client ?"""
    html = os.path.join(fol, "lettre.html")
    meta = _meta_path(fol)
    if not (os.path.exists(html) and os.path.exists(meta)):
        return False
    try:
        m = json.load(open(meta, encoding="utf-8"))
    except Exception:
        return False
    return m.get("hash") == client.get("hash")


def ensure_letter(client, root, template=None, logo_uri="", force=False, company=None):
    """Génère (ou réutilise le cache) la lettre d'un client. Retourne les chemins."""
    letters_dir = os.path.join(root, "Letters")
    os.makedirs(letters_dir, exist_ok=True)
    fol = os.path.join(letters_dir, folder_name(client))

    if not force and is_cached(client, fol):
        return _paths(client, fol, root, cached=True)

    os.makedirs(fol, exist_ok=True)
    if template is None:
        template = find_template(root)
    fill_html(client, os.path.join(fol, "lettre.html"), logo_uri, company)
    fill_docx(template, client, os.path.join(fol, "lettre.docx"))
    json.dump({"hash": client.get("hash"), "nom": client.get("nom_prenom")},
              open(_meta_path(fol), "w", encoding="utf-8"), ensure_ascii=False)
    return _paths(client, fol, root, cached=False)


def _paths(client, fol, root, cached):
    html = os.path.join(fol, "lettre.html")
    docx_p = os.path.join(fol, "lettre.docx")
    rel = lambda p: os.path.relpath(p, root) if os.path.exists(p) else ""
    return {
        "letter_dir": os.path.relpath(fol, root),
        "letter_html": rel(html),
        "letter_docx": rel(docx_p),
        "cached": cached,
    }


def logo_data_uri(custom=""):
    """Logo à utiliser : personnalisé (data URI) sinon logo Romana par défaut."""
    if custom and custom.startswith("data:image"):
        return custom
    import paths
    p = paths.resource("static", "logo.png")
    if not os.path.exists(p):
        return ""
    import base64
    return "data:image/png;base64," + base64.b64encode(open(p, "rb").read()).decode()


# --------------------------------------------------------------------------- #
#  Génération en lot : fichiers plats dans un dossier choisi                   #
# --------------------------------------------------------------------------- #

def _find_soffice():
    import shutil
    for c in ("soffice", "libreoffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if os.path.exists(c) or shutil.which(c):
            return c
    return None


def output_filename(client, pattern=""):
    """Nom de fichier plat pour un client (sans extension)."""
    nom = _safe(client.get("nom_prenom", "").split(" / ")[0]) or "CLIENT"
    cin = _safe((client.get("cin_str") or "").split(" / ")[0])
    ref = "-".join(x for x in (client.get("ilot", ""), client.get("immeuble", ""),
                               str(client.get("numero", ""))) if x)
    if pattern:
        name = pattern
        for k, val in (("{nom}", nom), ("{cin}", cin), ("{ilot}", client.get("ilot", "")),
                       ("{imm}", client.get("immeuble", "")), ("{no}", str(client.get("numero", ""))),
                       ("{ref}", ref), ("{type}", client.get("type_bien", ""))):
            name = name.replace(k, str(val))
        return _safe(name) or f"Lettre_{nom}"
    parts = ["Lettre", nom] + ([cin] if cin else []) + ([ref] if ref else [])
    return "_".join(parts)


def html_to_pdf(html_path, out_pdf):
    """Convertit un HTML en PDF hors-ligne. Essaie LibreOffice, puis wkhtmltopdf."""
    import subprocess, shutil, tempfile
    soffice = _find_soffice()
    out_dir = os.path.dirname(out_pdf)
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                            "--outdir", out_dir, html_path],
                           capture_output=True, timeout=90)
            produced = os.path.join(out_dir, os.path.splitext(os.path.basename(html_path))[0] + ".pdf")
            if os.path.exists(produced):
                if os.path.abspath(produced) != os.path.abspath(out_pdf):
                    shutil.move(produced, out_pdf)
                return True
        except Exception:
            pass
    if shutil.which("wkhtmltopdf"):
        try:
            subprocess.run(["wkhtmltopdf", "--quiet", html_path, out_pdf],
                           capture_output=True, timeout=90)
            if os.path.exists(out_pdf):
                return True
        except Exception:
            pass
    return False


def generate_batch(clients, out_dir, formats=("pdf",), date_override="",
                   company=None, naming="", logo_uri="", template=None,
                   progress=None):
    """Génère des fichiers plats (PDF et/ou DOCX) dans `out_dir`.

    Retourne une liste de résultats : {client, name, files:[...], warnings:[...]}.
    """
    os.makedirs(out_dir, exist_ok=True)
    if not logo_uri:
        logo_uri = logo_data_uri()
    results = []
    used = set()
    for i, c in enumerate(clients):
        if progress:
            progress(c.get("nom_prenom") or str(c.get("numero", "")),
                     int(100 * (i + 1) / max(1, len(clients))))
        base = output_filename(c, naming)
        name = base
        n = 2
        while name.lower() in used:
            name = f"{base}_{n}"; n += 1
        used.add(name.lower())

        files, warns = [], []
        want_pdf = "pdf" in formats
        want_docx = "docx" in formats

        html_tmp = os.path.join(out_dir, name + ".html")
        keep_html = False

        if want_pdf:
            fill_html(c, html_tmp, logo_uri, company, date_override)
            pdf_path = os.path.join(out_dir, name + ".pdf")
            if html_to_pdf(html_tmp, pdf_path):
                files.append(os.path.basename(pdf_path))
            else:
                # Pas de moteur PDF : on garde le HTML imprimable à la place.
                warns.append("PDF indisponible sur ce poste (installez LibreOffice). "
                             "Un fichier HTML imprimable est fourni.")
                files.append(name + ".html")
                keep_html = True

        if want_docx:
            docx_path = os.path.join(out_dir, name + ".docx")
            cc = dict(c)
            if date_override:
                cc["date_courrier"] = date_override
            if fill_docx(template, cc, docx_path):
                files.append(os.path.basename(docx_path))
            else:
                warns.append("DOCX indisponible (modèle de lettre introuvable).")

        # Nettoyer le HTML temporaire sauf si on doit le garder.
        if not keep_html and os.path.exists(html_tmp):
            try:
                os.remove(html_tmp)
            except OSError:
                pass

        results.append({"id": c.get("id"), "name": name,
                        "client": c.get("nom_prenom", ""), "files": files,
                        "warnings": warns})
    return results


def preview_html(client, logo_uri="", company=None, date_override=""):
    """HTML d'aperçu (sans bouton d'impression) pour l'affichage dans l'app."""
    if not logo_uri:
        logo_uri = logo_data_uri()
    return build_html(client, logo_uri, company, date_override, show_print=False)


def generate_all(clients, root, template=None, logo_uri="", force=False, progress=None):
    """Génère toutes les lettres (option « auto-générer »)."""
    if template is None:
        template = find_template(root)
    if not logo_uri:
        logo_uri = logo_data_uri()
    for i, c in enumerate(clients):
        if progress:
            progress(f"Lettre : {c.get('nom_prenom') or c['numero']}",
                     int(100 * (i + 1) / max(1, len(clients))))
        paths = ensure_letter(c, root, template, logo_uri, force=force)
        c.update(paths)
    return clients


if __name__ == "__main__":
    import sys
    root = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.dirname(HERE)
    data = json.load(open(os.path.join(HERE, "data.json"), encoding="utf-8"))
    generate_all(data["clients"], root, force="--force" in sys.argv,
                 progress=lambda m, p: print(f"[{p:3d}%] {m}"))
    json.dump(data, open(os.path.join(HERE, "data.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    print(f"[+] Lettres générées dans {os.path.join(root, 'Letters')}")
